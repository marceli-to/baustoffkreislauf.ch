import csv, sys, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

# usage: python unused-assets-xlsx.py <assets:unused csv> <out.xlsx>   (needs openpyxl)
src, out = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(src, encoding='utf-8')))
rows.sort(key=lambda r: (r['path'].rsplit('/', 1)[0] if '/' in r['path'] else '', r['path'].lower()))

F = 'Arial'
hdr_font = Font(name=F, bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='1D1D1B')
input_fill = PatternFill('solid', fgColor='FFFF00')
base = Font(name=F)
thin = Side(style='thin', color='D0D0D0')

wb = Workbook()
ov = wb.active
ov.title = 'Übersicht'
ws = wb.create_sheet('Dateien')

# Data sheet
headers = ['Ordner', 'Datei', 'Typ', 'Grösse (MB)', 'Geändert', 'Ansehen', 'Löschen?', 'Bemerkung']
ws.append(headers)
for c in ws[1]:
    c.font, c.fill = hdr_font, hdr_fill
    c.alignment = Alignment(vertical='center')
for r in rows:
    path = r['path']
    folder, name = (path.rsplit('/', 1) if '/' in path else ('(Hauptordner)', path))
    ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
    protected = folder.split('/')[0] == 'protected'
    url = 'https://baustoffkreislauf.ch/assets/' + urllib.parse.quote(path)
    ws.append([folder, name, ext, round(int(r['bytes']) / 1048576, 2), r['modified'],
               None, None, 'Geschützte Datei – nur nach Login abrufbar' if protected else None])
    row = ws.max_row
    link = ws.cell(row=row, column=6)
    if protected:
        link.value = '–'
    else:
        link.value, link.hyperlink = 'öffnen', url
        link.font = Font(name=F, color='0563C1', underline='single')
last = ws.max_row
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row:
        if c.column != 6 or c.value == '–':
            c.font = base
        c.border = Border(bottom=thin)
    row[3].number_format = '0.00'
    row[6].fill = input_fill
    row[6].alignment = Alignment(horizontal='center')

dv = DataValidation(type='list', formula1='"Ja,Nein"', allow_blank=True)
dv.error, dv.errorTitle = 'Bitte «Ja» oder «Nein» wählen.', 'Ungültiger Wert'
ws.add_data_validation(dv)
dv.add(f'G2:G{last}')
ws.conditional_formatting.add(f'A2:H{last}', FormulaRule(formula=['$G2="Ja"'], fill=PatternFill('solid', fgColor='F8D7DA')))
ws.conditional_formatting.add(f'A2:H{last}', FormulaRule(formula=['$G2="Nein"'], fill=PatternFill('solid', fgColor='D4EDDA')))

for col, w in zip('ABCDEFGH', [30, 60, 7, 12, 12, 10, 10, 45]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:H{last}'

# Overview sheet
D = f"Dateien!$D$2:$D${last}"
G = f"Dateien!$G$2:$G${last}"
A = f"Dateien!$A$2:$A${last}"
ov['A1'] = 'Nicht verwendete Dateien – baustoffkreislauf.ch'
ov['A1'].font = Font(name=F, bold=True, size=14)
ov['A2'] = 'Stand: Inhalte vom 22.10.2026. Auf diese Dateien verweist keine Seite, kein Beitrag, keine Publikation und kein Global der Website.'
ov['A3'] = 'Achtung: Dateien, die nur in Newslettern, E-Mails oder auf anderen Websites verlinkt sind, erscheinen hier trotzdem – bitte vor dem Löschen prüfen.'
ov['A5'] = 'So geht es'
ov['A5'].font = Font(name=F, bold=True)
ov['A6'] = '1. Im Blatt «Dateien» in der gelben Spalte «Löschen?» je Datei «Ja» oder «Nein» wählen. Leer lassen heisst: bleibt.'
ov['A7'] = '2. Mit «öffnen» lässt sich jede Datei ansehen. Bemerkungen in die letzte Spalte.'
ov['A8'] = '3. Datei zurückschicken – gelöscht werden nur die Zeilen mit «Ja».'
ov['A9'] = 'Beispiel: Löschen? = «Nein», Bemerkung = «Im Mitglieder-Newsletter März verlinkt»'
ov['A9'].font = Font(name=F, italic=True)

ov['A11'], ov['B11'], ov['C11'] = 'Zusammenfassung', 'Dateien', 'MB'
ov['A12'], ov['B12'], ov['C12'] = 'Total nicht verwendet', f'=COUNTA({A})', f'=SUM({D})'
ov['A13'], ov['B13'], ov['C13'] = 'Zum Löschen markiert («Ja»)', f'=COUNTIF({G},"Ja")', f'=SUMIF({G},"Ja",{D})'
ov['A14'], ov['B14'], ov['C14'] = 'Behalten («Nein»)', f'=COUNTIF({G},"Nein")', f'=SUMIF({G},"Nein",{D})'
ov['A15'], ov['B15'], ov['C15'] = 'Noch offen', '=B12-B13-B14', '=C12-C13-C14'

ov['A17'], ov['B17'], ov['C17'] = 'Nach Ordner', 'Dateien', 'MB'
folders = []
for r in rows:
    f = r['path'].split('/')[0] if '/' in r['path'] else '(Hauptordner)'
    if f not in folders:
        folders.append(f)
totals = {f: sum(int(x['bytes']) for x in rows if (x['path'].split('/')[0] if '/' in x['path'] else '(Hauptordner)') == f) for f in folders}
folders.sort(key=lambda f: -totals[f])
r0 = 18
for i, f in enumerate(folders):
    row = r0 + i
    ov.cell(row=row, column=1, value=f)
    crit = f if f == '(Hauptordner)' else f + '*'
    ov.cell(row=row, column=2, value=f'=COUNTIF({A},"{crit}")')
    ov.cell(row=row, column=3, value=f'=SUMIF({A},"{crit}",{D})')

for row in ov.iter_rows(min_row=1, max_row=r0 + len(folders)):
    for c in row:
        if c.font.name != F:
            c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, size=c.font.size)
        if c.font.name == F and not (c.font.bold or c.font.italic or c.font.size == 14):
            c.font = base
for cell in ('A11', 'B11', 'C11', 'A17', 'B17', 'C17'):
    ov[cell].font, ov[cell].fill = hdr_font, hdr_fill
ov['A1'].font = Font(name=F, bold=True, size=14)
ov['A5'].font = Font(name=F, bold=True)
ov['A9'].font = Font(name=F, italic=True)
for row in range(12, r0 + len(folders)):
    ov.cell(row=row, column=3).number_format = '0.0'
ov.column_dimensions['A'].width = 34
ov.column_dimensions['B'].width = 10
ov.column_dimensions['C'].width = 10
wb.calculation.fullCalcOnLoad = True
wb.save(out)
print(len(rows), 'rows,', len(folders), 'folders')
